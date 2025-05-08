import ast
import json
import logging
import os
import re
from typing import Dict, List, Tuple
import time
import pandas as pd
from openai import OpenAI
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Image, Spacer
import glob

class CISComplianceAnalyzer:
    def __init__(self, openai_key: str):
        self.client = OpenAI(api_key=openai_key)

    @staticmethod
    def chunk_json(data: Dict, max_tokens=1500) -> List[str]:
        json_str = json.dumps(data, indent=2)  # Optional: easier chunking and readable prompt
        lines = json_str.splitlines()
        chunks, chunk, token_count = [], [], 0

        for line in lines:
            line_tokens = len(line.split())
            if line.strip() == "":
                continue

            if token_count + line_tokens > max_tokens:
                if chunk:
                    chunks.append("\n".join(chunk))
                chunk, token_count = [line], line_tokens
            else:
                chunk.append(line)
                token_count += line_tokens

        if chunk:
            chunks.append("\n".join(chunk))

        return [c for c in chunks if c.strip()]

    @staticmethod
    def json_fix_and_load(response_str: str) -> Dict:
        try:
            match = re.search(r"\{.*}", response_str, re.DOTALL)
            if match:
                content = match.group(0)
                try:
                    return json.loads(content)
                except json.JSONDecodeError:
                    # Try replacing single quotes with double quotes
                    content_fixed = content.replace("'", "\"")
                    try:
                        return json.loads(content_fixed)
                    except json.JSONDecodeError:
                        # Last resort: literal_eval
                        return ast.literal_eval(content)

        except Exception as e:
            logging.warning(f"JSON parsing failed: {e}")
            logging.debug(f"Failed content: {response_str}")
        return {}

    def extract_actual_compliance(self, json_file: str) -> Tuple[
        Dict[str, str], Dict[str, str], Dict[str, Dict[str, str]], str]:
        if not os.path.exists(json_file):
            return {}, {}, {}, "❌ JSON file does not exist."

        try:
            with open(json_file, "r") as f:
                raw_data = json.load(f)
        except Exception as e:
            return {}, {}, {}, f"❌ Failed to load JSON: {str(e)}"

        chunks = self.chunk_json(raw_data, max_tokens=1500)
        actual_results, titles, resources = {}, {}, {}

        for i, chunk in enumerate(chunks):
            prompt = (
                "You are an expert in AWS compliance. Please analyze the following JSON data chunk and return it in this format:\n"
                "{\n  \"check_id\": {\"status\": \"PASS/FAIL\", \"resource_name\": \"...\", \"resource_type\": \"...\", \"check_title\": \"...\"}\n}\n\n"
                "Please analyze the following JSON chunk:\n"
                f"{chunk}\n\n"
                "Ensure the response is valid JSON formatted as per the example above."
            )

            try:
                response = self.client.chat.completions.create(
                    model="gpt-4o",
                    messages=[{"role": "system", "content": "You are a JSON parsing assistant."},
                              {"role": "user", "content": prompt}],
                    max_tokens=2000,
                    temperature=0.2
                )

                response_content = response.choices[0].message.content.strip()
                logging.debug(f"Raw API response for chunk {i + 1}: {response_content}")

                parsed_chunk = self.json_fix_and_load(response_content)
                if not parsed_chunk:
                    logging.warning(f"Chunk {i + 1} returned no parsable JSON. Response was: {response_content}")
                    continue

                for check_id, details in parsed_chunk.items():
                    actual_results[check_id] = details.get("status", "UNKNOWN").upper()
                    titles[check_id] = details.get("check_title", "UNKNOWN")
                    resources[check_id] = {
                        "resource_name": details.get("resource_name", "UNKNOWN"),
                        "resource_type": details.get("resource_type", "UNKNOWN")
                    }

            except Exception as e:
                logging.error(f"❌ Failed to extract compliance findings in chunk {i + 1}: {e}")
                continue

        if not actual_results:
            return {}, {}, {}, "❌ No valid compliance data found."

        return actual_results, titles, resources, "✅ Actual compliance data extracted."

    def analyze_expected_from_tf(self, tf_file: str, check_ids: List[str]) -> Dict[str, str]:
        if not os.path.exists(tf_file):
            return {check_id: "UNKNOWN" for check_id in check_ids}

        with open(tf_file, "r") as f:
            tf_content = f.read()

        prompt = (
            "You are an AWS CIS compliance and Terraform expert. Analyze the following Terraform code and evaluate ONLY the listed CIS check_ids.\n\n"
            "Instructions:\n"
            "- Only consider resources that are explicitly declared in the provided Terraform configuration.\n"
            "- If the exact setting related to the control is not present, mark it as \"FAIL\".\n"
            "- Do NOT assume defaults. Do NOT infer PASS based on related configurations.\n"
            "- Be strict. Do not return \"UNKNOWN\".\n"
            "- Output must be a JSON object, like:\n"
            "{\n  \"<check_id>\": \"PASS\" or \"FAIL\" \n}\n\n"
            f"Check IDs:\n{check_ids}\n\n"
            f"Terraform Code:\n{tf_content}"
        )

        try:
            response = self.client.chat.completions.create(
                model="gpt-4o",
                messages=[{"role": "system", "content": "You analyze Terraform for AWS compliance."},
                          {"role": "user", "content": prompt}],
                max_tokens=3000,
                temperature=0.2
            )

            result = self.json_fix_and_load(response.choices[0].message.content.strip())
            return {check_id: result.get(check_id, "FAIL") for check_id in check_ids}

        except Exception as e:
            logging.error(f"GPT Terraform analysis failed: {e}")
            return {check_id: "UNKNOWN" for check_id in check_ids}

    @staticmethod
    def add_footer(canvas, doc):
        # Footer disclaimer text with HTML-style tags
        disclaimer_text = (
            '<b>AI Disclaimer:</b> This report was generated using AI-based analysis. '
            'While the results are based on industry standards, they may contain inaccuracies. '
            'All configurations and findings should be manually validated against the official documentation '
            'of cloud providers and CIS Benchmarks for the compared analysis.'
        )

        # Create a Paragraph for styled text
        styles = getSampleStyleSheet()
        footer_style = styles["BodyText"]
        footer_style.fontSize = 8  # Adjust font size

        # Create the Paragraph object with the disclaimer text
        footer_paragraph = Paragraph(disclaimer_text, footer_style)

        # Get the bottom of the page position
        footer_height = 30  # You can adjust the position here if needed
        footer_paragraph.wrapOn(canvas, doc.pagesize[0] - 60, footer_height)
        footer_paragraph.drawOn(canvas, 30, footer_height)

    @staticmethod
    def xlsx_to_pdf(xlsx_path: str, pdf_path: str):
        df = pd.read_excel(xlsx_path)

        styles = getSampleStyleSheet()
        header_style = styles["Heading4"]
        body_style = styles["BodyText"]

        # Build table data
        data = [[Paragraph(str(col), header_style) for col in df.columns]]
        data += [[Paragraph(str(cell), body_style) for cell in row] for row in df.values]

        # Define uniform column widths
        num_cols = len(df.columns)
        total_width = 7.0 * inch
        col_widths = [total_width / num_cols] * num_cols
        # Table formatting
        table = Table(data, colWidths=col_widths, repeatRows=1)
        table.setStyle(TableStyle([
            ('GRID', (0, 0), (-1, -1), 0.5, colors.black),
            ('BACKGROUND', (0, 0), (-1, 0), colors.lightgrey),
            ('FONTSIZE', (0, 0), (-1, -1), 8),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 6),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ]))

        # Logo image
        logo = Image("logos/neova_solutions_logo.jpeg", width=1.5 * inch, height=1.5 * inch)
        logo.hAlign = 'RIGHT'

        # PDF generation
        pdf = SimpleDocTemplate(pdf_path, pagesize=A4)
        elements = [
            Table([[Paragraph("CIS Compliance Analysis Report", styles["Title"]), logo]], colWidths=[4.5 * inch, 2 * inch]),
            Spacer(1, 12),
            table
        ]
        pdf.build(elements, onFirstPage=CISComplianceAnalyzer.add_footer, onLaterPages=CISComplianceAnalyzer.add_footer)

    @staticmethod
    def generate_xlsx_report(actual: Dict[str, str], expected: Dict[str, str], titles: Dict[str, str],
                             resources: Dict[str, Dict[str, str]]) -> str:
        wb = Workbook()
        ws = wb.active
        ws.title = "CIS Compliance Analysis Report"

        headers = ["Control Name", "Resource Name", "Resource Type", "Actual Compliance", "Expected Compliance"]
        ws.append(headers)

        for col in "ABCDE":
            cell = ws[f"{col}1"]
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="DDDDDD", fill_type="solid")
            cell.alignment = Alignment(horizontal="center")

        for check_id in actual:
            ws.append([
                titles.get(check_id, "UNKNOWN"),
                resources.get(check_id, {}).get("resource_name", "UNKNOWN"),
                resources.get(check_id, {}).get("resource_type", "UNKNOWN"),
                actual.get(check_id, "UNKNOWN"),
                expected.get(check_id, "UNKNOWN"),
            ])

        for row in ws.iter_rows(min_row=2, min_col=4, max_col=5):
            for cell in row:
                if cell.value == "PASS":
                    cell.fill = PatternFill(start_color="C6EFCE", fill_type="solid")
                elif cell.value == "FAIL":
                    cell.fill = PatternFill(start_color="FFC7CE", fill_type="solid")
                else:
                    cell.fill = PatternFill(start_color="FFFFCC", fill_type="solid")

        
        timestamp = int(time.time())  # current epoch timestamp in seconds
        xlsx_path = f"neoComplianceAgent_compliance_report_{timestamp}.xlsx"
        pdf_path = f"neoComplianceAgent_compliance_report_{timestamp}.pdf"
        wb.save(xlsx_path)

        # Now call the fixed static method
        CISComplianceAnalyzer.xlsx_to_pdf(xlsx_path, pdf_path)

        return xlsx_path

    def fine_tune_results(self, actual: Dict[str, str], expected: Dict[str, str], tf_file: str = "main.tf") -> Tuple[Dict[str, str], str]:
        if not os.path.exists(tf_file):
            return expected, "❌ main.tf not found for fine-tuning."

        with open(tf_file, "r") as f:
            tf_content = f.read()

        prompt = (
            "You are an AWS Terraform compliance expert. You are given:\n"
            f"- Actual compliance results: {json.dumps(actual, indent=2)}\n"
            f"- Current expected results (from Terraform analysis): {json.dumps(expected, indent=2)}\n"
            "- The original Terraform configuration is shown below.\n\n"
            "Please revalidate the expected compliance results. If any 'expected' result is inaccurate or missed based on the actual findings and the Terraform code, correct it.\n"
            "Output only a corrected expected results map in the following format:\n"
            "{\n  \"check_id\": \"PASS/FAIL/UNKNOWN\"\n}"
            "\n\nTerraform Code:\n"
            f"{tf_content}"
        )

        try:
            response = self.client.chat.completions.create(
                model="gpt-4o",
                messages=[
                    {"role": "system", "content": "You are a Terraform compliance tuning assistant."},
                    {"role": "user", "content": prompt}
                ],
                max_tokens=2000,
                temperature=0.2
            )

            corrected = self.json_fix_and_load(response.choices[0].message.content.strip())
            updated_expected = expected.copy()
            updated_expected.update(corrected)

            return updated_expected, "✅ Fine-tuning of data completed and expected compliance corrected if necessary."

        except Exception as e:
            logging.error(f"Fine-tuning failed: {e}")
            return expected, "❌ Fine-tuning failed."

    def analyze_for_compliance_controls(self, json_file: str, tf_file: str = "main.tf") -> Tuple[str, str]:
        actual, titles, resources, msg = self.extract_actual_compliance(json_file)
        if not actual:
            return msg, ""

        expected = self.analyze_expected_from_tf(tf_file, list(actual.keys()))
        expected, tune_msg = self.fine_tune_results(actual, expected, tf_file)
        logging.info(tune_msg)

        xlsx_file = self.generate_xlsx_report(actual, expected, titles, resources)
        return "✅ Compliance analysis completed successfully.", xlsx_file

    def cleanup_report_files(self):
        """Cleans up report files with specific extensions."""
        extensions = ["xlsx", "pdf"]
        for ext in extensions:
            pattern = f"./neoComplianceAgent_compliance_report_*.{ext}"
            for report_file in glob.glob(pattern):
                try:
                    os.remove(report_file)
                    logging.info(f"🗑️ Deleted report file: {report_file}")
                except Exception as e:
                    logging.error(f"❌ Failed to delete {report_file}: {e}")